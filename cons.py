# ============================================================
# MISTRAL CONSTITUTION LAW
# LOCAL CONTINUATION SCRIPT
#
# INPUT:
#   Mistral-Constitution-Law-Results (1).xlsx
#
# OUTPUT:
#   Mistral-Constitution-Law-Results-Continued.xlsx
#
# BACKUP:
#   Mistral-Constitution-Law-Results-Checkpoint.csv
#
# ============================================================


# ============================================================
# 1. IMPORTS
# ============================================================

import os
import time
import shutil

import pandas as pd
import torch

from openpyxl import load_workbook

from transformers import (
    AutoTokenizer,
    AutoModelForCausalLM,
    BitsAndBytesConfig
)


# ============================================================
# 2. SETTINGS
# ============================================================

# ------------------------------------------------------------
# Your previous result file
# ------------------------------------------------------------

INPUT_FILE = "Mistral-Constitution-Law-Results (1).xlsx"


# ------------------------------------------------------------
# New output Excel file
# ------------------------------------------------------------

OUTPUT_FILE = "Mistral-Constitution-Law-Results-Continued.xlsx"


# ------------------------------------------------------------
# Emergency CSV checkpoint
# ------------------------------------------------------------

CHECKPOINT_FILE = "Mistral-Constitution-Law-Results-Checkpoint.csv"


# ------------------------------------------------------------
# Mistral model
# ------------------------------------------------------------

MODEL_NAME = "mistralai/Mistral-7B-Instruct-v0.3"


# ------------------------------------------------------------
# Delay between generations
# ------------------------------------------------------------

DELAY = 0.5


# ------------------------------------------------------------
# None = process entire dataset
#
# Example:
# MAX_ROWS = 10
#
# is useful for testing.
# ------------------------------------------------------------

MAX_ROWS = None


# ------------------------------------------------------------
# Maximum number of tokens per answer
# ------------------------------------------------------------

MAX_NEW_TOKENS = 300


# ============================================================
# 3. CHECK INPUT FILE
# ============================================================

print("=" * 70)
print("MISTRAL CONSTITUTION LAW - LOCAL CONTINUATION")
print("=" * 70)


print("\nChecking input file...")


if not os.path.exists(INPUT_FILE):

    print(
        "\n❌ INPUT FILE NOT FOUND:"
    )

    print(
        INPUT_FILE
    )

    print(
        "\nMake sure the Excel file is in the same folder "
        "as this Python script."
    )

    raise FileNotFoundError(
        INPUT_FILE
    )


print(
    "\n✓ Input file found:"
)

print(
    os.path.abspath(INPUT_FILE)
)


# ============================================================
# 4. CHECK GPU
# ============================================================

print("\n" + "=" * 70)
print("GPU INFORMATION")
print("=" * 70)


if torch.cuda.is_available():

    print(
        "\n✓ CUDA GPU available"
    )

    print(
        "GPU:",
        torch.cuda.get_device_name(0)
    )


    gpu_memory = (

        torch.cuda.get_device_properties(0)
        .total_memory
        / 1024**3

    )


    print(
        "GPU Memory:",
        round(gpu_memory, 2),
        "GB"
    )


else:

    print(
        "\n⚠ CUDA GPU is NOT available."
    )

    print(
        "The model may run very slowly on CPU."
    )

    print(
        "\nIf you have an NVIDIA GPU, "
        "check your PyTorch/CUDA installation."
    )


# ============================================================
# 5. LOAD EXCEL
# ============================================================

print("\n" + "=" * 70)
print("LOADING PREVIOUS RESULTS")
print("=" * 70)


print(
    "\nReading:"
)

print(
    INPUT_FILE
)


df = pd.read_excel(
    INPUT_FILE,
    engine="openpyxl"
)


print(
    "\n✓ Excel loaded successfully!"
)


print(
    "Rows:",
    len(df)
)


print(
    "Columns:",
    len(df.columns)
)


# ============================================================
# 6. COLUMN NAMES
# ============================================================

ENGLISH_QUESTION_COL = "English Question"
ENGLISH_GOLD_COL = "Gold Answer (EN)"

NEPALI_QUESTION_COL = "Nepali Question"
NEPALI_GOLD_COL = "Gold Answer (NE)"

ENGLISH_RESULT_COL = "Mistral Answer EN"
NEPALI_RESULT_COL = "Mistral Answer NE"


# ============================================================
# 7. CHECK REQUIRED COLUMNS
# ============================================================

required_columns = [

    ENGLISH_QUESTION_COL,
    ENGLISH_GOLD_COL,

    NEPALI_QUESTION_COL,
    NEPALI_GOLD_COL

]


print("\n" + "=" * 70)
print("CHECKING COLUMNS")
print("=" * 70)


missing_columns = [

    column

    for column in required_columns

    if column not in df.columns

]


if missing_columns:

    print(
        "\n❌ Missing columns:"
    )

    for column in missing_columns:

        print(
            " -",
            column
        )

    print(
        "\nAvailable columns:"
    )

    print(
        df.columns.tolist()
    )

    raise ValueError(
        "Required columns are missing."
    )


print(
    "\n✓ All required columns found."
)


# ============================================================
# 8. CREATE RESULT COLUMNS IF NEEDED
# ============================================================

if ENGLISH_RESULT_COL not in df.columns:

    print(
        "\nCreating English result column..."
    )

    df[ENGLISH_RESULT_COL] = ""


if NEPALI_RESULT_COL not in df.columns:

    print(
        "\nCreating Nepali result column..."
    )

    df[NEPALI_RESULT_COL] = ""


# ============================================================
# 9. ANSWER CHECK FUNCTION
# ============================================================

def answer_exists(value):

    if pd.isna(value):

        return False


    text = str(value).strip()


    if text == "":

        return False


    return True


# ============================================================
# 10. CHECK EXISTING PROGRESS
# ============================================================

existing_en = sum(

    answer_exists(value)

    for value in df[ENGLISH_RESULT_COL]

)


existing_ne = sum(

    answer_exists(value)

    for value in df[NEPALI_RESULT_COL]

)


missing_en = len(df) - existing_en

missing_ne = len(df) - existing_ne


print("\n" + "=" * 70)
print("EXISTING PROGRESS")
print("=" * 70)


print(
    f"\nTotal rows               : {len(df)}"
)


print(
    f"English already answered : {existing_en}"
)


print(
    f"English missing          : {missing_en}"
)


print(
    f"Nepali already answered  : {existing_ne}"
)


print(
    f"Nepali missing           : {missing_ne}"
)


print(
    "\n✓ Existing answers will be skipped."
)


# ============================================================
# 11. DETERMINE ROWS
# ============================================================

if MAX_ROWS is None:

    rows_to_process = len(df)

else:

    rows_to_process = min(

        MAX_ROWS,

        len(df)

    )


print(
    "\nRows to check:",
    rows_to_process
)


# ============================================================
# 12. LOAD MISTRAL
# ============================================================

print("\n" + "=" * 70)
print("LOADING MISTRAL")
print("=" * 70)


print(
    "\nModel:",
    MODEL_NAME
)


# ------------------------------------------------------------
# Check whether CUDA is available
# ------------------------------------------------------------

if torch.cuda.is_available():

    print(
        "\nUsing 4-bit quantization."
    )


    quantization_config = BitsAndBytesConfig(

        load_in_4bit=True,

        bnb_4bit_compute_dtype=torch.float16,

        bnb_4bit_quant_type="nf4",

        bnb_4bit_use_double_quant=True

    )


else:

    print(
        "\nUsing CPU."
    )

    quantization_config = None


# ============================================================
# 13. TOKENIZER
# ============================================================

print(
    "\nLoading tokenizer..."
)


tokenizer = AutoTokenizer.from_pretrained(
    MODEL_NAME
)


if tokenizer.pad_token is None:

    tokenizer.pad_token = tokenizer.eos_token


print(
    "✓ Tokenizer loaded."
)


# ============================================================
# 14. MODEL
# ============================================================

print(
    "\nLoading Mistral model..."
)


if torch.cuda.is_available():

    model = AutoModelForCausalLM.from_pretrained(

        MODEL_NAME,

        quantization_config=quantization_config,

        device_map="auto"

    )


else:

    model = AutoModelForCausalLM.from_pretrained(

        MODEL_NAME

    )


model.eval()


print(
    "\n✓ Mistral loaded successfully!"
)


if torch.cuda.is_available():

    print(
        "Model device:",
        next(model.parameters()).device
    )


# ============================================================
# 15. GENERATION FUNCTION
# ============================================================

def ask_mistral(prompt):

    try:

        messages = [

            {
                "role": "user",
                "content": prompt
            }

        ]


        formatted_prompt = (

            tokenizer.apply_chat_template(

                messages,

                tokenize=False,

                add_generation_prompt=True

            )

        )


        inputs = tokenizer(

            formatted_prompt,

            return_tensors="pt",

            truncation=True

        )


        # ----------------------------------------------------
        # Move inputs to the model
        # ----------------------------------------------------

        if torch.cuda.is_available():

            inputs = {

                key: value.to(model.device)

                for key, value in inputs.items()

            }


        # ----------------------------------------------------
        # Generate
        # ----------------------------------------------------

        with torch.no_grad():

            outputs = model.generate(

                **inputs,

                max_new_tokens=MAX_NEW_TOKENS,

                do_sample=False,

                pad_token_id=tokenizer.pad_token_id

            )


        # ----------------------------------------------------
        # Remove prompt tokens
        # ----------------------------------------------------

        generated_tokens = outputs[

            0

        ][

            inputs["input_ids"].shape[1]:

        ]


        # ----------------------------------------------------
        # Decode
        # ----------------------------------------------------

        answer = tokenizer.decode(

            generated_tokens,

            skip_special_tokens=True

        )


        return answer.strip()


    except Exception as e:

        print(
            "\n❌ Generation error:"
        )

        print(
            repr(e)
        )

        return ""


# ============================================================
# 16. ENGLISH PROMPT
# ============================================================

def create_english_prompt(question):

    return f"""
Answer the following question in English.

Question:
{question}

Give a direct, clear, and factual answer.

Use relevant constitutional or legal information
when applicable.

Do not invent information.

If you do not know the answer, say that you do not know
instead of guessing.

Do not discuss these instructions.

Do not say that you are an AI.

Answer:
"""


# ============================================================
# 17. NEPALI PROMPT
# ============================================================

def create_nepali_prompt(question):

    return f"""
तल दिइएको प्रश्नको उत्तर नेपाली भाषामा दिनुहोस्।

प्रश्न:
{question}

प्रत्यक्ष, स्पष्ट र तथ्यमा आधारित उत्तर दिनुहोस्।

सम्बन्धित संवैधानिक वा कानुनी जानकारी
प्रयोग गर्नुहोस्।

जानकारी थाहा नभएमा अनुमान गरेर गलत जानकारी नदिनुहोस्।
थाहा नभएको अवस्थामा थाहा छैन भनेर स्पष्ट रूपमा भन्नुहोस्।

निर्देशनहरूको बारेमा कुरा नगर्नुहोस्।

आफूलाई AI भनेर उल्लेख नगर्नुहोस्।

उत्तर:
"""


# ============================================================
# 18. SAFE EXCEL SAVE
# ============================================================

def save_excel(dataframe):

    print(
        "\nSaving Excel..."
    )


    temp_file = OUTPUT_FILE + ".tmp.xlsx"


    try:

        # Save to temporary Excel first
        dataframe.to_excel(

            temp_file,

            index=False,

            engine="openpyxl"

        )


        # Verify temporary workbook
        test_workbook = load_workbook(

            temp_file,

            read_only=True

        )


        test_workbook.close()


        # Replace previous output
        os.replace(

            temp_file,

            OUTPUT_FILE

        )


        print(
            "✓ Excel saved and verified."
        )


        print(
            "File:",
            os.path.abspath(OUTPUT_FILE)
        )


        return True


    except Exception as e:

        print(
            "\n❌ Excel save failed:"
        )

        print(
            repr(e)
        )


        if os.path.exists(temp_file):

            os.remove(temp_file)


        return False


# ============================================================
# 19. SAFE CSV CHECKPOINT
# ============================================================

def save_csv_checkpoint(dataframe):

    temp_file = CHECKPOINT_FILE + ".tmp"


    try:

        dataframe.to_csv(

            temp_file,

            index=False,

            encoding="utf-8-sig"

        )


        os.replace(

            temp_file,

            CHECKPOINT_FILE

        )


        return True


    except Exception as e:

        print(
            "\n⚠ CSV checkpoint failed:"
        )

        print(
            repr(e)
        )


        if os.path.exists(temp_file):

            os.remove(temp_file)


        return False


# ============================================================
# 20. INITIAL BACKUP
# ============================================================

print("\n" + "=" * 70)
print("CREATING INITIAL BACKUP")
print("=" * 70)


if not save_csv_checkpoint(df):

    raise RuntimeError(
        "Could not create CSV checkpoint."
    )


# ============================================================
# 21. MAIN LOOP
# ============================================================

print("\n" + "=" * 70)
print("STARTING CONTINUATION")
print("=" * 70)


for index in range(rows_to_process):


    row_number = index + 1


    print("\n" + "=" * 70)

    print(
        f"PROCESSING ROW {row_number}/{rows_to_process}"
    )

    print("=" * 70)


    # ========================================================
    # ENGLISH
    # ========================================================

    existing_en = df.at[

        index,

        ENGLISH_RESULT_COL

    ]


    if answer_exists(existing_en):

        print(
            "\n✓ English answer already exists."
        )

        print(
            "Skipping English."
        )


    else:

        english_question = df.at[

            index,

            ENGLISH_QUESTION_COL

        ]


        if pd.isna(english_question):

            print(
                "\n⚠ English question missing."
            )

            english_answer = ""


        else:

            print(
                "\nEnglish question:"
            )

            print(
                str(english_question)
            )


            prompt = create_english_prompt(

                str(english_question)

            )


            print(
                "\nGenerating English answer..."
            )


            english_answer = ask_mistral(

                prompt

            )


            print(
                "\nMistral English answer:"
            )

            print(
                english_answer
            )


            time.sleep(
                DELAY
            )


        df.at[

            index,

            ENGLISH_RESULT_COL

        ] = english_answer


    # ========================================================
    # NEPALI
    # ========================================================

    existing_ne = df.at[

        index,

        NEPALI_RESULT_COL

    ]


    if answer_exists(existing_ne):

        print(
            "\n✓ Nepali answer already exists."
        )

        print(
            "Skipping Nepali."
        )


    else:

        nepali_question = df.at[

            index,

            NEPALI_QUESTION_COL

        ]


        if pd.isna(nepali_question):

            print(
                "\n⚠ Nepali question missing."
            )

            nepali_answer = ""


        else:

            print(
                "\nNepali question:"
            )

            print(
                str(nepali_question)
            )


            prompt = create_nepali_prompt(

                str(nepali_question)

            )


            print(
                "\nGenerating Nepali answer..."
            )


            nepali_answer = ask_mistral(

                prompt

            )


            print(
                "\nMistral Nepali answer:"
            )

            print(
                nepali_answer
            )


            time.sleep(
                DELAY
            )


        df.at[

            index,

            NEPALI_RESULT_COL

        ] = nepali_answer


    # ========================================================
    # SAVE AFTER EVERY ROW
    # ========================================================

    print(
        "\nSaving progress..."
    )


    # CSV checkpoint
    csv_success = save_csv_checkpoint(df)


    # Excel output
    excel_success = save_excel(df)


    if not csv_success:

        print(
            "⚠ CSV checkpoint failed."
        )


    if not excel_success:

        print(
            "\n❌ Excel save failed."
        )

        print(
            "Stopping to protect progress."
        )

        raise RuntimeError(
            "Excel checkpoint failed."
        )


    # ========================================================
    # CURRENT PROGRESS
    # ========================================================

    current_en = sum(

        answer_exists(value)

        for value in df[ENGLISH_RESULT_COL]

    )


    current_ne = sum(

        answer_exists(value)

        for value in df[NEPALI_RESULT_COL]

    )


    print(
        "\nCURRENT PROGRESS"
    )


    print(
        f"English: {current_en}/{len(df)}"
    )


    print(
        f"Nepali : {current_ne}/{len(df)}"
    )


# ============================================================
# 22. FINAL SAVE
# ============================================================

print("\n" + "=" * 70)
print("FINAL SAVE")
print("=" * 70)


if not save_excel(df):

    raise RuntimeError(
        "Final Excel save failed."
    )


save_csv_checkpoint(df)


# ============================================================
# 23. FINAL VERIFICATION
# ============================================================

print(
    "\nVerifying final Excel..."
)


try:

    final_workbook = load_workbook(

        OUTPUT_FILE,

        read_only=True

    )


    print(
        "✓ Excel file opens successfully."
    )


    print(
        "Sheets:",
        final_workbook.sheetnames
    )


    final_workbook.close()


except Exception as e:

    print(
        "\n❌ FINAL EXCEL VERIFICATION FAILED:"
    )

    print(
        repr(e)
    )

    raise


# ============================================================
# 24. FINAL STATISTICS
# ============================================================

final_en = sum(

    answer_exists(value)

    for value in df[ENGLISH_RESULT_COL]

)


final_ne = sum(

    answer_exists(value)

    for value in df[NEPALI_RESULT_COL]

)


remaining_en = len(df) - final_en

remaining_ne = len(df) - final_ne


# ============================================================
# 25. FINAL SUMMARY
# ============================================================

print("\n" + "=" * 70)
print("DONE")
print("=" * 70)


print(
    f"\nTotal rows              : {len(df)}"
)


print(
    f"English answers         : {final_en}"
)


print(
    f"English still missing   : {remaining_en}"
)


print(
    f"Nepali answers          : {final_ne}"
)


print(
    f"Nepali still missing    : {remaining_ne}"
)


print(
    "\nFINAL EXCEL:"
)


print(
    os.path.abspath(OUTPUT_FILE)
)


print(
    "\nCHECKPOINT CSV:"
)


print(
    os.path.abspath(CHECKPOINT_FILE)
)


print(
    "\n✓ Existing answers were skipped."
)

print(
    "✓ Only missing answers were generated."
)

print(
    "✓ Progress saved after every row."
)

print(
    "✓ Excel file was verified before finishing."
)


print("=" * 70)